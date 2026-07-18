import io
import os
import shutil
from pathlib import Path

from datetime import datetime, timedelta
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from modules.url_utils import get_brand_key


# 內容通常很長的欄位，固定給較寬的欄寬並自動換行，
# 不要用「配合最長內容」的方式撐爆整欄。
LONG_TEXT_COLUMN_WIDTH = 50
LONG_TEXT_COLUMNS = {
    "AI判斷原因",
    "商品內容",
    "評論",
    "台灣代理證據",
    "展覽商品分類",
    "連絡人資料",
    "後續連絡情況",
}

# 內容通常很短、固定格式的欄位，給窄一點的欄寬。
NARROW_COLUMN_WIDTHS = {
    "編號": 8,
    "記錄日期": 12,
    "搜尋名稱": 20,
    "資料來源": 14,
    "AI分類": 12,
    "是否適合代理": 12,
    "代理推薦分數": 12,
    "台灣代理狀態": 14,
    "台灣檢查信心分數": 14,
    "展覽年份": 10,
    "人工確認": 12,
}

MIN_COLUMN_WIDTH = 10
MAX_AUTO_COLUMN_WIDTH = 30

# 部分欄位內容可能非常長（例如展覽官方名錄的完整分類標籤，
# 未過濾的話可能高達數百字），自動換行後會把整列撐得異常高，
# 這裡限制每列最多顯示的行數，超出的部分仍保留在儲存格內，
# 使用者可以手動放大該列或點開儲存格查看完整內容。
DEFAULT_LINE_HEIGHT_POINTS = 15
MAX_WRAPPED_LINES = 8

# 每次寫入正式資料庫前，先留一份備份，
# 避免寫入中斷、或事後發現這次結果有誤時無法還原。
# 只保留最近幾份，避免備份資料夾無限增大。
BACKUP_DIR_NAME = "backups"
MAX_BACKUPS_TO_KEEP = 20


def set_worksheet_column_widths(
    worksheet,
    dataframe: pd.DataFrame,
) -> None:
    """
    依欄位內容調整合理欄寬，避免用預設固定寬度
    導致長文字被截斷、短欄位又太寬不好閱讀。

    共用邏輯，供「寫入既有 Excel 檔案」與
    「匯出成記憶體中的 Excel 內容」共同呼叫。
    """

    wrap_alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    for column_index, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        column_letter = get_column_letter(
            column_index
        )

        if column_name in LONG_TEXT_COLUMNS:
            worksheet.column_dimensions[
                column_letter
            ].width = LONG_TEXT_COLUMN_WIDTH

            for cell in worksheet[column_letter][
                1:
            ]:
                cell.alignment = wrap_alignment

            continue

        if column_name in NARROW_COLUMN_WIDTHS:
            worksheet.column_dimensions[
                column_letter
            ].width = NARROW_COLUMN_WIDTHS[
                column_name
            ]

            continue

        longest_value_length = max(
            [
                len(str(column_name))
            ]
            + [
                len(str(value))
                for value in dataframe[
                    column_name
                ]
                if value not in ("", None)
            ]
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                longest_value_length + 2,
                MIN_COLUMN_WIDTH,
            ),
            MAX_AUTO_COLUMN_WIDTH,
        )

    cap_wrapped_row_heights(
        worksheet,
        dataframe,
    )


def cap_wrapped_row_heights(
    worksheet,
    dataframe: pd.DataFrame,
) -> None:
    """
    限制自動換行欄位撐出的列高上限。

    列高本來會依最長的換行欄位自動決定，
    但像展覽官方名錄的分類標籤這類欄位內容
    可能長達數百字，換算下來會有十幾行，
    導致單一列異常地高、蓋掉版面。
    這裡估算每列需要幾行，並限制在
    MAX_WRAPPED_LINES 以內，內容仍完整保留在
    儲存格裡，只是預設顯示不會全部展開。
    """

    long_text_column_letters = [
        get_column_letter(
            dataframe.columns.get_loc(column_name)
            + 1
        )
        for column_name in LONG_TEXT_COLUMNS
        if column_name in dataframe.columns
    ]

    if not long_text_column_letters:
        return

    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        max_lines = 1

        for column_letter in long_text_column_letters:
            value = worksheet[
                f"{column_letter}{row_index}"
            ].value

            if not value:
                continue

            column_width = (
                worksheet.column_dimensions[
                    column_letter
                ].width
                or LONG_TEXT_COLUMN_WIDTH
            )

            estimated_lines = -(
                -len(str(value))
                // max(
                    int(column_width),
                    1,
                )
            )

            max_lines = max(
                max_lines,
                estimated_lines,
            )

        capped_lines = min(
            max_lines,
            MAX_WRAPPED_LINES,
        )

        worksheet.row_dimensions[
            row_index
        ].height = (
            capped_lines
            * DEFAULT_LINE_HEIGHT_POINTS
        )


def apply_column_widths(
    output_path: str,
    dataframe: pd.DataFrame,
    sheet_name: str = "總表",
) -> None:
    """
    依欄位內容調整既有 Excel 檔案的欄寬，直接存回原檔案。
    """

    workbook = load_workbook(output_path)
    worksheet = workbook[sheet_name]

    set_worksheet_column_widths(
        worksheet,
        dataframe,
    )

    workbook.save(output_path)


def backup_existing_file(output_path: str) -> None:
    """
    寫入正式檔案前，把目前版本備份一份到 backups/ 資料夾，
    檔名帶時間戳，只保留最近 MAX_BACKUPS_TO_KEEP 份。

    output_path 尚不存在時（第一次執行）不需要備份。
    """
    path = Path(output_path)

    if not path.exists():
        return

    backup_dir = path.parent / BACKUP_DIR_NAME
    backup_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"{path.stem}_{timestamp}{path.suffix}"

    shutil.copy2(path, backup_path)

    existing_backups = sorted(
        backup_dir.glob(f"{path.stem}_*{path.suffix}"),
        key=lambda backup: backup.stat().st_mtime,
    )

    for old_backup in existing_backups[:-MAX_BACKUPS_TO_KEEP]:
        old_backup.unlink()


def write_dataframe_atomically(
    dataframe: pd.DataFrame,
    output_path: str,
    sheet_name: str = "總表",
) -> None:
    """
    先把完整內容（含欄寬調整）寫到暫存檔，
    確認全部寫完沒有出錯後，才用系統的原子替換動作
    換成正式檔案。

    這樣不管搜尋執行到一半被中斷、斷電，或正式檔案
    剛好被其他程式佔用，任何時間點看到的 output_path
    都只會是「完整寫完的舊版本」或「完整寫完的新版本」，
    不會有寫一半、損毀的中間狀態。
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    tmp_path = path.with_name(
        f"{path.stem}.tmp{path.suffix}"
    )

    dataframe.to_excel(
        tmp_path,
        index=False,
        sheet_name=sheet_name,
    )

    try:
        apply_column_widths(
            str(tmp_path),
            dataframe,
            sheet_name=sheet_name,
        )

    except Exception as error:
        print(
            "[COLUMN WIDTH ADJUST FAILED] "
            f"{error}"
        )

    backup_existing_file(output_path)

    os.replace(tmp_path, path)


def export_dataframe_to_excel_bytes(
    dataframe: pd.DataFrame,
    sheet_name: str = "篩選結果",
) -> bytes:
    """
    把 DataFrame 匯出成 Excel 內容（bytes），
    欄寬套用跟主資料庫一樣的自動調整邏輯，
    供 Streamlit 下載按鈕直接使用，不需要先寫入磁碟。
    """

    raw_buffer = io.BytesIO()

    dataframe.to_excel(
        raw_buffer,
        index=False,
        sheet_name=sheet_name,
    )

    raw_buffer.seek(0)

    workbook = load_workbook(raw_buffer)
    worksheet = workbook[sheet_name]

    set_worksheet_column_widths(
        worksheet,
        dataframe,
    )

    output_buffer = io.BytesIO()
    workbook.save(output_buffer)

    return output_buffer.getvalue()


COLUMNS = [
    "記錄日期",
    "搜尋名稱",
    "編號",
    "公司名稱",
    "網站",
    "國家",
    "資料來源",
    "展覽名稱",
    "展覽年份",
    "展覽攤位",
    "展覽公司名稱",
    "展覽商品分類",
    "展覽參展類型",
    "展覽來源頁面",
    "商品類別",
    "商品內容",
    "AI分類",
    "是否適合代理",
    "代理推薦分數",
    "AI判斷原因",
    "台灣代理狀態",
    "台灣代理商名稱",
    "台灣代理證據",
    "台灣代理來源",
    "台灣檢查信心分數",
    "評論",
    "人工確認",
    "後續連絡情況",
    "連絡人資料",
    "來源連結",
]

# 只要「人工確認」欄位有填任何內容（例如「是」、「V」），
# 代表這筆資料已經有人手動核對／修正過，之後重新分析
# （到期或強制重新分析）時整筆跳過，不覆蓋任何欄位——
# 不像 MANUAL_COLUMNS 只保護特定幾欄，這是保護整筆紀錄。
LOCK_COLUMN = "人工確認"



# 這些欄位可能由公司人員人工填寫。
# 程式更新品牌資料時，不應把人工內容覆蓋掉。
MANUAL_COLUMNS = [
    "後續連絡情況",
    "連絡人資料",
]




def normalize_brand_name(name: str) -> str:
    """
    品牌名稱標準化，作為網域缺失時的備用去重方式。
    """
    if not isinstance(name, str):
        return ""

    return "".join(
        character.lower()
        for character in name.strip()
        if character.isalnum()
    )


def make_brand_key(record: dict) -> str:
    """
    優先使用官方網域建立品牌識別鍵。
    沒有網址時才使用公司名稱。
    """
    website = record.get("網站", "")
    source_url = record.get("來源連結", "")

    domain = get_brand_key(website) or get_brand_key(source_url)

    if domain:
        return f"domain:{domain}"

    brand_name = normalize_brand_name(record.get("公司名稱", ""))

    if brand_name:
        return f"name:{brand_name}"

    return ""


def prepare_dataframe(records: list[dict]) -> pd.DataFrame:
    """
    統一 DataFrame 欄位。
    """
    dataframe = pd.DataFrame(records)

    for column in COLUMNS:
        if column not in dataframe.columns:
            dataframe[column] = ""

    return dataframe[COLUMNS]


def load_existing_records(output_path: str) -> list[dict]:
    """
    讀取既有 Excel 品牌資料。

    Excel 不存在、內容為空或讀取失敗時，回傳空清單。
    """
    path = Path(output_path)

    if not path.exists():
        return []

    try:
        dataframe = pd.read_excel(
            path,
            sheet_name="總表",
            dtype=object,
        )

        dataframe = dataframe.fillna("")

        return dataframe.to_dict(orient="records")

    except Exception as exc:
        print(f"[EXCEL READ FAILED] {output_path}: {exc}")
        return []

def parse_record_date(value) -> datetime | None:
    """
    解析 Excel 內的記錄日期。

    支援：
    20260620
    2026-06-20
    Excel datetime
    """
    if value in ("", None):
        return None

    if isinstance(value, datetime):
        return value

    text = str(value).strip()

    # Excel 可能把整數讀成 20260620.0
    if text.endswith(".0"):
        text = text[:-2]

    date_formats = [
        "%Y%m%d",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y-%m-%d %H:%M:%S",
    ]

    for date_format in date_formats:
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            continue

    return None


def load_recent_existing_domains(
        output_path: str,
        refresh_days: int = 30,
) -> set[str]:
    """
    讀取 Excel 內近期已分析品牌的網域。

    只有在 refresh_days 內的品牌才會被跳過；
    過期品牌仍可重新分析。
    """
    existing_records = load_existing_records(output_path)

    if not existing_records:
        return set()

    cutoff_date = datetime.now() - timedelta(days=refresh_days)
    recent_domains: set[str] = set()

    for record in existing_records:
        website = record.get("網站", "")
        source_url = record.get("來源連結", "")

        domain = (
            get_brand_key(website)
            or get_brand_key(source_url)
        )

        if not domain:
            continue

        record_date = parse_record_date(
            record.get("記錄日期", "")
        )

        # 舊資料若沒有可解析日期，先視為需要重新分析，
        # 避免錯誤地永久跳過。
        if record_date is None:
            continue

        if record_date >= cutoff_date:
            recent_domains.add(domain)

    return recent_domains
def is_locked_record(record: dict) -> bool:
    """
    判斷這筆資料是否已被人工確認、鎖定不再覆蓋。

    只要「人工確認」欄位有填任何內容（不限格式），
    就視為鎖定。
    """
    return bool(
        str(record.get(LOCK_COLUMN, "")).strip()
    )


def merge_vendor_records(
        existing_records: list[dict],
        new_records: list[dict],
) -> tuple[list[dict], int, int, int]:
    """
    合併舊資料與新資料。

    回傳：
    - 合併後資料
    - 新增筆數
    - 更新筆數
    - 鎖定跳過筆數（已人工確認，本次完全不覆蓋）
    """
    merged_by_key: dict[str, dict] = {}
    key_order: list[str] = []

    # 先放入既有資料。
    for index, record in enumerate(existing_records):
        record_copy = dict(record)
        key = make_brand_key(record_copy)

        # 無法建立識別鍵時，仍保留資料。
        if not key:
            key = f"existing_unknown:{index}"

        if key not in merged_by_key:
            key_order.append(key)

        merged_by_key[key] = record_copy

    added_count = 0
    updated_count = 0
    locked_count = 0

    # 再合併本次新結果。
    for index, new_record in enumerate(new_records):
        new_record_copy = dict(new_record)
        key = make_brand_key(new_record_copy)

        if not key:
            key = f"new_unknown:{index}"

        if key in merged_by_key:
            old_record = merged_by_key[key]

            # 已人工確認的資料整筆跳過，不覆蓋任何欄位，
            # 保留人工修正過的內容（不限於 MANUAL_COLUMNS）。
            if is_locked_record(old_record):
                locked_count += 1
                continue

            # 先保留人工欄位。
            manual_values = {
                column: old_record.get(column, "")
                for column in MANUAL_COLUMNS
            }

            # 自動分析欄位以新結果為準。
            old_record.update(new_record_copy)

            # 將人工欄位放回去，避免被空白覆蓋。
            for column, old_value in manual_values.items():
                new_value = new_record_copy.get(column, "")

                if old_value not in ("", None):
                    old_record[column] = old_value
                else:
                    old_record[column] = new_value

            merged_by_key[key] = old_record
            updated_count += 1

        else:
            merged_by_key[key] = new_record_copy
            key_order.append(key)
            added_count += 1

    merged_records = [
        merged_by_key[key]
        for key in key_order
    ]

    # 重新排列編號。
    for index, record in enumerate(merged_records, start=1):
        record["編號"] = index

    return (
        merged_records,
        added_count,
        updated_count,
        locked_count,
    )


def export_vendor_records(
    records: list[dict],
    output_path: str,
    incremental: bool = True,
) -> dict:
    """
    匯出品牌資料。

    incremental=True：
    讀取舊 Excel，合併後再輸出。

    incremental=False：
    僅輸出本次搜尋結果。
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    existing_records = (
        load_existing_records(output_path)
        if incremental
        else []
    )

    (
        merged_records,
        added_count,
        updated_count,
        locked_count,
    ) = merge_vendor_records(
        existing_records=existing_records,
        new_records=records,
    )

    dataframe = prepare_dataframe(merged_records)

    write_dataframe_atomically(
        dataframe,
        output_path,
    )

    summary = {
        "本次找到": len(records),
        "新增品牌": added_count,
        "更新品牌": updated_count,
        "鎖定跳過": locked_count,
        "資料庫總數": len(merged_records),
    }

    print(
        "[INCREMENTAL EXPORT] "
        f"本次找到 {summary['本次找到']} 筆，"
        f"新增 {summary['新增品牌']} 筆，"
        f"更新 {summary['更新品牌']} 筆，"
        f"鎖定跳過 {summary['鎖定跳過']} 筆，"
        f"總數 {summary['資料庫總數']} 筆"
    )

    return summary